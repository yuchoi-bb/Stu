"""문서 입력 파이프라인 (§8): PDF/DOCX/TXT/MD → 텍스트 추출.

- Bedrock에는 텍스트만 전달. 대용량은 청킹.
- HWP는 미지원 (§12.1 확인 항목) — 명시적으로 안내 메시지 반환.
- 추출 실패는 예외 대신 안내 텍스트로 반환 (파이프라인 중단 방지).
"""
import os
import zipfile

from . import config, db, logs

_log = logs.get("docparse")

CHUNK_CHARS = 40_000   # 청크당 문자 수 (대용량 문서 분할)

# 자원 고갈 방어(§8): worker 1 프로세스라 추출 중 OOM은 전체 장애 → 상한을 둔다.
MAX_EXTRACT_CHARS = 2_000_000            # 추출 텍스트 상한(~2MB) — 초과분은 절단
MAX_PDF_PAGES = 1_000                    # PDF 페이지 상한
MAX_DECOMPRESSED_BYTES = 200 * 1024 * 1024   # DOCX(zip) 압축 해제 총량 상한
MAX_COMPRESSION_RATIO = 200              # 압축비 상한 (압축 폭탄 탐지)
_TRUNC = "\n[... 추출 상한 도달, 이후 생략]"


def extract_text(path: str, mime_type: str | None, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    try:
        if ext in (".txt", ".md") or (mime_type or "").startswith("text/"):
            return _read_text(path)
        if ext == ".pdf" or mime_type == "application/pdf":
            return _read_pdf(path)
        if ext == ".docx" or (mime_type or "").endswith("wordprocessingml.document"):
            return _read_docx(path)
        if ext in (".hwp", ".hwpx"):
            return f"[HWP 미지원: {filename} — 추출 도구 미도입(§12.1). 텍스트로 변환 후 재첨부 필요]"
        return f"[미지원 포맷: {filename} ({ext or mime_type}) — 텍스트 추출 생략]"
    except Exception as e:
        return f"[추출 실패: {filename} — {e}]"


def _read_text(path: str) -> str:
    with open(path, encoding="utf-8", errors="replace") as f:
        data = f.read(MAX_EXTRACT_CHARS + 1)
    return data[:MAX_EXTRACT_CHARS] + _TRUNC if len(data) > MAX_EXTRACT_CHARS else data


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    n = len(reader.pages)
    parts, total = [], 0
    for i in range(min(n, MAX_PDF_PAGES)):
        t = reader.pages[i].extract_text() or ""
        if t.strip():
            parts.append(f"--- page {i + 1} ---\n{t}")
            total += len(t)
            if total > MAX_EXTRACT_CHARS:
                parts.append(_TRUNC)
                break
    if n > MAX_PDF_PAGES:
        parts.append(f"[... {n - MAX_PDF_PAGES}개 페이지 생략(상한 {MAX_PDF_PAGES})]")
    return "\n\n".join(parts) if parts else "[PDF에서 추출된 텍스트 없음 (스캔 이미지일 수 있음)]"


def _docx_bomb_check(path: str) -> None:
    """DOCX(zip) 압축 해제 총량/압축비가 과다하면 압축 폭탄으로 보고 거부."""
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
    except zipfile.BadZipFile:
        return   # zip이 아니면 docx.Document가 깔끔히 실패 처리
    total = sum(i.file_size for i in infos)
    comp = sum(i.compress_size for i in infos) or 1
    if total > MAX_DECOMPRESSED_BYTES or total / comp > MAX_COMPRESSION_RATIO:
        _log.warning("docx 압축 폭탄 의심: total=%d ratio=%.1f", total, total / comp)
        raise ValueError(f"압축 해제 크기 과다({total // (1024 * 1024)}MB) — 처리 거부")


def _read_docx(path: str) -> str:
    _docx_bomb_check(path)
    import docx
    d = docx.Document(path)
    parts, total = [], 0
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
            total += len(p.text)
            if total > MAX_EXTRACT_CHARS:
                parts.append(_TRUNC)
                return "\n".join(parts)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                line = " | ".join(cells)
                parts.append(line)
                total += len(line)
                if total > MAX_EXTRACT_CHARS:
                    parts.append(_TRUNC)
                    return "\n".join(parts)
    return "\n".join(parts) if parts else "[DOCX에서 추출된 텍스트 없음]"


def chunks(text: str) -> list[str]:
    """대용량 문서를 CHUNK_CHARS 단위로 분할 (문단 경계 우선)."""
    if len(text) <= CHUNK_CHARS:
        return [text]
    out, buf = [], []
    size = 0
    for para in text.split("\n"):
        if size + len(para) > CHUNK_CHARS and buf:
            out.append("\n".join(buf))
            buf, size = [], 0
        buf.append(para)
        size += len(para) + 1
    if buf:
        out.append("\n".join(buf))
    return out


def process_attachment(attachment_id: int) -> str:
    """업로드 직후 호출 (ThreadPool). 추출 → parsed_text_path 저장."""
    row = db.one("SELECT * FROM attachments WHERE attachment_id=?", (attachment_id,))
    if row is None:
        return ""
    text = extract_text(row["storage_path"], row["mime_type"], row["filename"])
    parsed_path = row["storage_path"] + ".txt"
    with open(parsed_path, "w", encoding="utf-8") as f:
        f.write(text)
    db.execute("UPDATE attachments SET parsed_text_path=? WHERE attachment_id=?",
               (parsed_path, attachment_id))
    return text


def session_attachment_text(session_id: str, *, max_chars: int = 120_000) -> str:
    """세션 첨부들의 추출 텍스트를 모아 생성 컨텍스트로 (§7.1 Step 2/3)."""
    rows = db.query(
        "SELECT filename, parsed_text_path FROM attachments "
        "WHERE session_id=? AND parsed_text_path IS NOT NULL ORDER BY attachment_id",
        (session_id,))
    blocks, total = [], 0
    for r in rows:
        try:
            with open(r["parsed_text_path"], encoding="utf-8") as f:
                content = f.read()
        except OSError:
            continue
        header = f"# 첨부: {r['filename']}\n"
        remaining = max_chars - total
        if remaining <= len(header):
            break
        body = content[:remaining - len(header)]
        blocks.append(header + body)
        total += len(header) + len(body)
    return "\n\n".join(blocks)
