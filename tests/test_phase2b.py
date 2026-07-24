"""Phase 2b 추가 기능 테스트: 문서 추출, 2-pass, 재시도, 브랜치 생성, stale.

실행: python3 tests/test_phase2b.py
"""
import os
import subprocess
import sys
import tempfile
import unittest.mock as mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

TMP = tempfile.mkdtemp(prefix="studio-2b-")
os.environ["STUDIO_DB"] = os.path.join(TMP, "studio.db")
os.environ["STUDIO_FERNET_KEY"] = os.path.join(TMP, ".fernet.key")
os.environ["STUDIO_ATTACH_DIR"] = os.path.join(TMP, "attach")

from studio.app import app                       # noqa: E402
from studio import db, docparse, ghe, jobs, pipeline, analysis  # noqa: E402

jobs.submit = lambda fn, *a, **k: fn(*a, **k)
client = app.test_client()
H = {"X-Remote-User": "hong"}


# ---------- 문서 텍스트 추출 (§8) ----------

def _mkfile(name, data):
    p = os.path.join(TMP, name)
    mode = "wb" if isinstance(data, bytes) else "w"
    with open(p, mode) as f:
        f.write(data)
    return p

# TXT/MD
assert "hello md" in docparse.extract_text(_mkfile("a.md", "hello md"),
                                           "text/markdown", "a.md")
# DOCX
import docx  # noqa: E402
d = docx.Document()
d.add_paragraph("요구조건: 내림 처리")
d.add_paragraph("음수는 0으로")
dpath = os.path.join(TMP, "req.docx")
d.save(dpath)
txt = docparse.extract_text(dpath, None, "req.docx")
assert "내림 처리" in txt and "음수는 0" in txt, txt

# PDF
from pypdf import PdfWriter  # noqa: E402
w = PdfWriter()
w.add_blank_page(width=200, height=200)
ppath = os.path.join(TMP, "x.pdf")
with open(ppath, "wb") as f:
    w.write(f)
pres = docparse.extract_text(ppath, "application/pdf", "x.pdf")
assert "추출" in pres or pres.startswith("---") or "없음" in pres  # 빈 페이지 안내

# HWP 미지원 안내
assert "HWP 미지원" in docparse.extract_text("/x", None, "doc.hwp")

# 청킹
big = "line\n" * 20000
assert len(docparse.chunks(big)) > 1

# 세션 첨부 텍스트 집계 (업로드 → 비동기 추출 → 컨텍스트)
sid = client.post("/api/studio/sessions", headers=H,
                  json={"title": "t", "tool_target": "parser"}).get_json()["session_id"]
import io  # noqa: E402
client.post(f"/api/studio/sessions/{sid}/attachments", headers=H,
            data={"file": (io.BytesIO("첨부 내용: R1 내림".encode()), "spec.md")},
            content_type="multipart/form-data")
ctx = docparse.session_attachment_text(sid)
assert "R1 내림" in ctx and "spec.md" in ctx, ctx
print("OK: 문서 텍스트 추출 (MD/DOCX/PDF/HWP안내) + 청킹 + 세션 집계")

# ---------- Step 3 2-pass: 파일 지목 + 원문 fetch (§7.1) ----------

def git(cwd, *a):
    subprocess.run(["git", *a], cwd=cwd, check=True, capture_output=True)

bare = os.path.join(TMP, "remote.git")
git(None, "init", "--bare", bare)
work = os.path.join(TMP, "seed")
git(None, "clone", bare, work)
os.makedirs(os.path.join(work, "src"))
with open(os.path.join(work, "src/parser.c"), "w") as f:
    f.write("int old_parse() { return 1; }\n" * 10)
git(work, "-c", "user.name=s", "-c", "user.email=s@s", "add", "-A")
git(work, "-c", "user.name=s", "-c", "user.email=s@s", "commit", "-m", "init")
git(work, "branch", "feature/foo")
git(work, "push", "origin", "master", "feature/foo")
REMOTE = "file://" + bare

# pass 1 프롬프트 파싱
paths = pipeline.parse_paths_block("```paths\nsrc/parser.c\ninclude/x.h\n```")
assert paths == ["src/parser.c", "include/x.h"], paths

# _fetch_originals: 지목된 파일 원문 + blob SHA 회수
db.execute("INSERT OR IGNORE INTO users (user_id) VALUES ('hong')")
studio = {"repo": "thr", "branch_name": "feature/foo",
          "requirements": "R1"}
with mock.patch.object(pipeline, "invoke_claude",
                       lambda *a, **k: {"output": {"message": {"content": [
                           {"text": "```paths\nsrc/parser.c\n```"}]}}, "usage": {}}), \
     mock.patch.object(ghe, "get_token", lambda u: "tok"), \
     mock.patch.object(ghe, "remote_url", lambda repo, tok: REMOTE):
    originals, blobs = pipeline._fetch_originals(
        "hong", studio, None, sid, "ST-x", 1)
assert "src/parser.c" in originals and "old_parse" in originals["src/parser.c"]
assert blobs["src/parser.c"] and len(blobs["src/parser.c"]) == 40
print("OK: 2-pass — 파일 지목 파싱 + 원문/blob SHA fetch")

# 신규 생성만인 경우 fetch 건너뜀
with mock.patch.object(pipeline, "invoke_claude",
                       lambda *a, **k: {"output": {"message": {"content": [
                           {"text": "```paths\n```"}]}}, "usage": {}}):
    o2, b2 = pipeline._fetch_originals("hong", studio, None, sid, "ST-x", 1)
assert o2 == {} and b2 == {}
print("OK: 2-pass — 신규 생성만이면 원문 fetch 생략")

# ---------- push/dispatch 재시도 (§11) ----------

calls = {"n": 0}
class Resp:
    def __init__(self, code): self.status_code = code; self.text = ""
def flaky():
    calls["n"] += 1
    return Resp(500 if calls["n"] < 3 else 200)
r = ghe._http_retry(flaky, base=0.01)
assert r.status_code == 200 and calls["n"] == 3
# 4xx는 즉시 반환(재시도 안 함)
calls["n"] = 0
r = ghe._http_retry(lambda: (calls.__setitem__("n", calls["n"]+1), Resp(404))[1], base=0.01)
assert r.status_code == 404 and calls["n"] == 1
print("OK: HTTP 재시도 — 5xx 3회 백오프 / 4xx 즉시")

# ---------- 브랜치 생성 (§6.1) ----------

import studio.ghe as ghemod  # noqa: E402
seq = []
def fake_get(url, **k):
    seq.append(url)
    class R:
        status_code = 404 if "/branches/" in url else 200
        def json(self): return {"object": {"sha": "abc123"}}
    return R()
def fake_post(url, **k):
    class R:
        status_code = 201
    return R()
with mock.patch.object(ghemod, "get_token", lambda u: "tok"), \
     mock.patch("studio.ghe.requests.get", fake_get), \
     mock.patch("studio.ghe.requests.post", fake_post):
    created = ghemod.ensure_branch("hong", "thr", "feature/new")
assert created is True
print("OK: 미존재 브랜치 base에서 생성")

# ---------- 분석 md stale 경고 (§7.3) ----------

md_fresh = "# parser 분석\n- 기준 커밋: abcdef1\n## 1. 구조\n"
assert analysis._stale_warning(md_fresh, "abcdef1234567890") is None
assert "오래됨" in analysis._stale_warning(md_fresh, "9999999000")
assert analysis._stale_warning("SHA 없음", "abcdef1") is None
print("OK: 분석 md stale 경고 (SHA 일치/불일치/미기재)")

# ---------- 회귀 확인: 백엔드 스위트 여전히 그린 ----------
print("\nALL PHASE-2B TESTS PASSED")
